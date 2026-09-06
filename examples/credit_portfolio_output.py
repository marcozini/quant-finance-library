#### Credit Portfolio Model Output ####

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

from quant_finance.credit.portfolio_model import (
    run_credit_portfolio_model,
)

from credit_portfolio_report import (
    create_credit_portfolio_report,
)


FILE_PATH = "data/credit_risk_input.xlsx"

OUTPUT_DIR = Path("outputs")

OUTPUT_EXCEL = (
    OUTPUT_DIR
    / "credit_portfolio_results.xlsx"
)

OUTPUT_PDF = (
    OUTPUT_DIR
    / "credit_portfolio_report.pdf"
)

LOSS_DISTRIBUTION_PLOT = (
    OUTPUT_DIR
    / "credit_loss_distribution.png"
)

TAIL_DISTRIBUTION_PLOT = (
    OUTPUT_DIR
    / "credit_loss_distribution_tail.png"
)

EXPOSURE_PLOT = (
    OUTPUT_DIR
    / "credit_exposure_concentration.png"
)

RISK_CONTRIBUTION_PLOT = (
    OUTPUT_DIR
    / "credit_risk_contributions.png"
)

PD_PLOT = (
    OUTPUT_DIR
    / "credit_pd_overview.png"
)

PLOT_FILES = [
    LOSS_DISTRIBUTION_PLOT,
    TAIL_DISTRIBUTION_PLOT,
    EXPOSURE_PLOT,
    RISK_CONTRIBUTION_PLOT,
    PD_PLOT,
]


# Read a boolean output setting from the central steering file.
def get_output_flag(
    settings,
    setting_name,
):

    value = settings[
        setting_name
    ]

    if isinstance(
        value,
        bool,
    ):
        return value

    if isinstance(
        value,
        (
            int,
            float,
        ),
    ):
        return bool(
            value
        )

    if isinstance(
        value,
        str,
    ):

        value = (
            value
            .strip()
            .lower()
        )

        if value in {
            "true",
            "yes",
            "1",
        }:
            return True

        if value in {
            "false",
            "no",
            "0",
        }:
            return False

    raise ValueError(
        f"Model setting '{setting_name}' "
        f"must be TRUE or FALSE."
    )


# Plot the complete simulated portfolio loss distribution.
def create_loss_distribution_plot(
    portfolio_losses,
    summary,
):

    mean_loss = summary[
        "migration_mean_loss"
    ]

    var = summary[
        "value_at_risk"
    ]

    es = summary[
        "expected_shortfall"
    ]

    confidence = summary[
        "confidence_level"
    ]

    base_currency = summary[
        "base_currency"
    ]

    plt.figure(
        figsize=(
            10,
            6,
        )
    )

    plt.hist(
        portfolio_losses,
        bins=100,
        alpha=0.8,
    )

    plt.axvline(
        mean_loss,
        linestyle="--",
        label=(
            f"Mean loss: "
            f"{mean_loss:.2f}"
        ),
    )

    plt.axvline(
        var,
        linestyle="--",
        label=(
            f"{confidence:.1%} VaR: "
            f"{var:.2f}"
        ),
    )

    plt.axvline(
        es,
        linestyle=":",
        label=(
            f"{confidence:.1%} ES: "
            f"{es:.2f}"
        ),
    )

    plt.title(
        "Simulated Credit Portfolio Loss Distribution"
    )

    plt.xlabel(
        f"Portfolio loss "
        f"({base_currency})"
    )

    plt.ylabel(
        "Number of simulations"
    )

    plt.legend()

    plt.tight_layout()

    plt.savefig(
        LOSS_DISTRIBUTION_PLOT,
        dpi=200,
    )

    plt.close()


# Plot the tail of the simulated portfolio loss distribution.
def create_tail_loss_distribution_plot(
    portfolio_losses,
    portfolio,
    summary,
):

    var = summary[
        "value_at_risk"
    ]

    es = summary[
        "expected_shortfall"
    ]

    confidence = summary[
        "confidence_level"
    ]

    base_currency = summary[
        "base_currency"
    ]

    default_losses = (
        portfolio[
            "exposure_base"
        ]
        * portfolio[
            "lgd"
        ]
    )

    largest_index = (
        default_losses
        .idxmax()
    )

    largest_default_loss = float(
        default_losses.loc[
            largest_index
        ]
    )

    largest_counterparty = (
        portfolio.loc[
            largest_index,
            "counterparty",
        ]
    )

    tail_start = max(
        0,
        0.6 * var,
    )

    tail_end = max(
        largest_default_loss
        * 1.25,
        es * 1.50,
    )

    tail_end = min(
        tail_end,
        float(
            np.max(
                portfolio_losses
            )
        ),
    )

    tail_losses = (
        portfolio_losses[
            (
                portfolio_losses
                >= tail_start
            )
            & (
                portfolio_losses
                <= tail_end
            )
        ]
    )

    plt.figure(
        figsize=(
            10,
            6,
        )
    )

    plt.hist(
        tail_losses,
        bins=80,
        alpha=0.8,
    )

    plt.axvline(
        var,
        linestyle="--",
        label=(
            f"{confidence:.1%} VaR: "
            f"{var:.2f}"
        ),
    )

    plt.axvline(
        es,
        linestyle=":",
        label=(
            f"{confidence:.1%} ES: "
            f"{es:.2f}"
        ),
    )

    plt.axvline(
        largest_default_loss,
        linestyle="-.",
        label=(
            f"{largest_counterparty} "
            f"default loss: "
            f"{largest_default_loss:.2f}"
        ),
    )

    plt.xlim(
        tail_start,
        tail_end,
    )

    plt.title(
        "Simulated Credit Portfolio Loss Distribution "
        "— Tail View"
    )

    plt.xlabel(
        f"Portfolio loss "
        f"({base_currency})"
    )

    plt.ylabel(
        "Number of simulations"
    )

    plt.legend()

    plt.tight_layout()

    plt.savefig(
        TAIL_DISTRIBUTION_PLOT,
        dpi=200,
    )

    plt.close()


# Plot portfolio exposure concentration in base currency.
def create_exposure_plot(
    portfolio,
    base_currency,
):

    plot_data = (
        portfolio
        .sort_values(
            "exposure_base",
            ascending=False,
        )
    )

    plt.figure(
        figsize=(
            10,
            6,
        )
    )

    plt.bar(
        plot_data[
            "counterparty"
        ],
        plot_data[
            "exposure_base"
        ],
    )

    plt.title(
        "Counterparty Exposure Concentration"
    )

    plt.xlabel(
        "Counterparty"
    )

    plt.ylabel(
        f"Exposure "
        f"({base_currency})"
    )

    plt.xticks(
        rotation=45,
        ha="right",
    )

    plt.tight_layout()

    plt.savefig(
        EXPOSURE_PLOT,
        dpi=200,
    )

    plt.close()


# Compare incremental VaR and Expected Shortfall.
def create_risk_contribution_plot(
    portfolio,
    confidence,
    base_currency,
):

    plot_data = (
        portfolio
        .sort_values(
            "incremental_es",
            ascending=False,
        )
    )

    positions = np.arange(
        len(
            plot_data
        )
    )

    width = 0.38

    plt.figure(
        figsize=(
            11,
            6,
        )
    )

    plt.bar(
        positions
        - width / 2,
        plot_data[
            "incremental_var"
        ],
        width,
        label=(
            f"Incremental "
            f"{confidence:.1%} VaR"
        ),
    )

    plt.bar(
        positions
        + width / 2,
        plot_data[
            "incremental_es"
        ],
        width,
        label=(
            f"Incremental "
            f"{confidence:.1%} ES"
        ),
    )

    plt.title(
        "Counterparty Tail-Risk Contributions"
    )

    plt.xlabel(
        "Counterparty"
    )

    plt.ylabel(
        f"Risk contribution "
        f"({base_currency})"
    )

    plt.xticks(
        positions,
        plot_data[
            "counterparty"
        ],
        rotation=45,
        ha="right",
    )

    plt.legend()

    plt.tight_layout()

    plt.savefig(
        RISK_CONTRIBUTION_PLOT,
        dpi=200,
    )

    plt.close()


# Plot counterparty probabilities of default.
def create_pd_plot(
    portfolio,
):

    plot_data = (
        portfolio
        .sort_values(
            "pd",
            ascending=False,
        )
        .copy()
    )

    # Avoid zero values on the logarithmic axis.
    plot_data[
        "pd_plot"
    ] = (
        plot_data[
            "pd"
        ]
        .clip(
            lower=1e-15
        )
    )

    plt.figure(
        figsize=(
            10,
            6,
        )
    )

    plt.bar(
        plot_data[
            "counterparty"
        ],
        plot_data[
            "pd_plot"
        ],
    )

    plt.yscale(
        "log"
    )

    plt.title(
        "Counterparty Probability of Default"
    )

    plt.xlabel(
        "Counterparty"
    )

    plt.ylabel(
        "One-year PD (log scale)"
    )

    plt.xticks(
        rotation=45,
        ha="right",
    )

    plt.tight_layout()

    plt.savefig(
        PD_PLOT,
        dpi=200,
    )

    plt.close()


# Create all charts required for Excel, PDF
# and optional PNG output.
def create_plots(
    portfolio_losses,
    portfolio,
    summary,
):

    confidence = summary[
        "confidence_level"
    ]

    base_currency = summary[
        "base_currency"
    ]

    create_loss_distribution_plot(
        portfolio_losses,
        summary,
    )

    create_tail_loss_distribution_plot(
        portfolio_losses,
        portfolio,
        summary,
    )

    create_exposure_plot(
        portfolio,
        base_currency,
    )

    create_risk_contribution_plot(
        portfolio,
        confidence,
        base_currency,
    )

    create_pd_plot(
        portfolio,
    )


# Format worksheet headers.
def format_sheet_headers(
    sheet,
):

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in sheet[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = (
            header_fill
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    sheet.freeze_panes = (
        "A2"
    )


# Automatically size worksheet columns.
def autosize_columns(
    sheet,
    max_width=32,
):

    for column in sheet.columns:

        letter = (
            column[
                0
            ]
            .column_letter
        )

        max_length = max(
            (
                len(
                    str(
                        cell.value
                    )
                )
                if cell.value
                is not None
                else 0
            )
            for cell
            in column[
                :100
            ]
        )

        sheet.column_dimensions[
            letter
        ].width = min(
            max_length + 2,
            max_width,
        )


# Format selected columns based on their header names.
def format_columns_by_header(
    sheet,
    headers,
    number_format,
):

    header_map = {
        cell.value:
            cell.column
        for cell
        in sheet[1]
    }

    for header in headers:

        if (
            header
            not in header_map
        ):
            continue

        column_index = (
            header_map[
                header
            ]
        )

        for row in range(
            2,
            sheet.max_row + 1,
        ):

            sheet.cell(
                row=row,
                column=column_index,
            ).number_format = (
                number_format
            )


# Format final Excel workbook and embed plots.
def format_excel_output():

    workbook = load_workbook(
        OUTPUT_EXCEL
    )

    for sheet_name in (
        workbook.sheetnames
    ):

        sheet = workbook[
            sheet_name
        ]

        format_sheet_headers(
            sheet
        )

        autosize_columns(
            sheet
        )

    # Portfolio summary dashboard.
    summary_sheet = workbook[
        "portfolio_summary"
    ]

    summary_sheet.column_dimensions[
        "A"
    ].width = 34

    summary_sheet.column_dimensions[
        "B"
    ].width = 22

    loss_image = Image(
        LOSS_DISTRIBUTION_PLOT
    )

    loss_image.width = 650
    loss_image.height = 390

    summary_sheet.add_image(
        loss_image,
        "D2",
    )

    tail_image = Image(
        TAIL_DISTRIBUTION_PLOT
    )

    tail_image.width = 650
    tail_image.height = 390

    summary_sheet.add_image(
        tail_image,
        "N2",
    )

    exposure_image = Image(
        EXPOSURE_PLOT
    )

    exposure_image.width = 650
    exposure_image.height = 390

    summary_sheet.add_image(
        exposure_image,
        "D24",
    )

    risk_image = Image(
        RISK_CONTRIBUTION_PLOT
    )

    risk_image.width = 650
    risk_image.height = 390

    summary_sheet.add_image(
        risk_image,
        "N24",
    )

    pd_image = Image(
        PD_PLOT
    )

    pd_image.width = 650
    pd_image.height = 390

    summary_sheet.add_image(
        pd_image,
        "D46",
    )

    # Counterparty results formatting.
    counterparty_sheet = workbook[
        "counterparty_results"
    ]

    format_columns_by_header(
        counterparty_sheet,
        [
            "exposure_local",
            "exposure_base",
            "expected_loss",
            "equity_value",
            "debt",
            "asset_value",
            "incremental_var",
            "incremental_es",
            "standalone_var",
            "standalone_es",
            "var_diversification_benefit",
            "es_diversification_benefit",
        ],
        "#,##0.0000",
    )

    format_columns_by_header(
        counterparty_sheet,
        [
            "lgd",
            "equity_volatility",
            "asset_volatility",
            "risk_free_rate",
            "coupon_rate",
        ],
        "0.0000%",
    )

    format_columns_by_header(
        counterparty_sheet,
        [
            "pd",
        ],
        "0.0000E+00",
    )

    format_columns_by_header(
        counterparty_sheet,
        [
            "fx_rate_to_base",
            "global_loading",
            "region_loading",
            "sector_loading",
            "idiosyncratic_loading",
            "marginal_var",
            "marginal_es",
        ],
        "0.000000",
    )

    # Model notes formatting.
    notes_sheet = workbook[
        "model_notes"
    ]

    notes_sheet.column_dimensions[
        "A"
    ].width = 30

    notes_sheet.column_dimensions[
        "B"
    ].width = 100

    for row in notes_sheet.iter_rows(
        min_row=2,
        min_col=2,
        max_col=2,
    ):

        row[
            0
        ].alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    workbook.save(
        OUTPUT_EXCEL
    )


# Create the detailed Excel model output.
def create_excel_output(
    portfolio,
    settings,
    factor_proxies,
    rating_migration_matrix,
    market_data,
    portfolio_losses,
    summary,
):

    confidence = summary[
        "confidence_level"
    ]

    base_currency = summary[
        "base_currency"
    ]

    # Portfolio summary.
    summary_df = pd.DataFrame({

        "metric": [
            "Base currency",
            "Number of counterparties",
            "Total exposure",
            "Default expected loss",
            "Migration mean loss",
            f"{confidence:.1%} VaR",
            (
                f"{confidence:.1%} "
                f"Expected Shortfall"
            ),
            "Unexpected loss",
            "Maximum simulated loss",
            "Minimum simulated loss",
        ],

        "value": [
            base_currency,
            len(
                portfolio
            ),
            portfolio[
                "exposure_base"
            ].sum(),
            summary[
                "default_expected_loss"
            ],
            summary[
                "migration_mean_loss"
            ],
            summary[
                "value_at_risk"
            ],
            summary[
                "expected_shortfall"
            ],
            summary[
                "unexpected_loss"
            ],
            summary[
                "maximum_simulated_loss"
            ],
            summary[
                "minimum_simulated_loss"
            ],
        ],
    })

    # Counterparty-level results.
    counterparty_columns = [
        "counterparty",
        "ticker",
        "factor_loading_proxy_ticker",

        "sector",
        "region",

        "pd_method",
        "rating",

        "exposure_local",
        "currency",
        "fx_rate_to_base",
        "fx_rate_source",
        "exposure_base",

        "lgd",

        "pd",
        "pd_source",
        "expected_loss",

        "equity_value",
        "equity_volatility",
        "debt",
        "asset_value",
        "asset_volatility",

        "risk_free_rate",
        "risk_free_rate_source",

        "maturity",
        "coupon_rate",
        "payment_frequency",

        "factor_calibration_ticker",
        "factor_source",

        "global_loading",
        "region_loading",
        "sector_loading",
        "idiosyncratic_loading",

        "incremental_var",
        "incremental_es",

        "marginal_var",
        "marginal_es",

        "standalone_var",
        "standalone_es",

        "var_diversification_benefit",
        "es_diversification_benefit",
    ]

    counterparty_results = (
        portfolio[
            counterparty_columns
        ]
        .copy()
    )

    counterparty_results.insert(
        counterparty_results
        .columns
        .get_loc(
            "exposure_base"
        )
        + 1,
        "base_currency",
        base_currency,
    )

    # Model settings.
    settings_df = pd.DataFrame(
        settings.items(),
        columns=[
            "setting",
            "value",
        ],
    )

    # Model notes and key assumptions.
    model_notes = pd.DataFrame({

        "topic": [
            "Portfolio inputs",
            "Merton model",
            "Manual Merton example",
            "Rating inputs",
            "FX conversion",
            "Risk-free rates",
            "Factor proxy fallback",
            "Dependence model",
            "Value at Risk",
            "Expected Shortfall",
            "LGD",
        ],

        "description": [

            (
                "Counterparty exposures, ratings, LGDs, bond terms and "
                "other portfolio positions are demonstration inputs."
            ),

            (
                "Merton PDs are structural model outputs rather than agency "
                "ratings. Very low PDs for financially strong listed companies "
                "illustrate an important limitation of the pure structural model."
            ),

            (
                "The private-company example uses synthetic manual equity value, "
                "equity volatility and debt inputs to demonstrate Merton "
                "calibration without listed-company market data."
            ),

            (
                "Rating assignments and the rating migration matrix are synthetic "
                "demonstration inputs and do not represent ratings from a rating agency."
            ),

            (
                "Exposures are converted into the selected portfolio base currency "
                "using automatically retrieved FX rates. When a specific market data "
                "date is selected, the latest available observation on or before that "
                "date is used. The frozen market data snapshot is used only as fallback. "
                "FX risk itself is not jointly simulated with credit risk."
            ),

            (
                "Risk-free rates are retrieved automatically by currency. When a "
                "specific market data date is selected, the latest available observation "
                "on or before that date is used. Stored flat policy-rate proxies are "
                "fallback inputs only and are not calibrated market yield curves."
            ),

            (
                "The factor-loading proxy ticker is used directly for unlisted "
                "counterparties and acts as fallback if factor estimation using a "
                "listed counterparty's own ticker fails. It does not replace company "
                "data used for Merton PD estimation."
            ),

            (
                "Counterparty dependence is modeled through common global, regional "
                "and sector factors together with idiosyncratic risk."
            ),

            (
                "VaR measures the loss quantile at the selected confidence level "
                "and may understate the severity of rare losses beyond that quantile."
            ),

            (
                "Expected Shortfall averages losses beyond VaR and therefore captures "
                "the severity of extreme tail scenarios and large single-counterparty "
                "concentrations more directly."
            ),

            (
                "LGD is deterministic and fixed at 50% in the demonstration portfolio "
                "to isolate exposure, credit quality and dependence effects."
            ),
        ],
    })

    # Scenario-level simulated losses.
    scenario_losses = pd.DataFrame({

        "scenario": range(
            1,
            len(
                portfolio_losses
            )
            + 1,
        ),

        "portfolio_loss":
            portfolio_losses,
    })

    with pd.ExcelWriter(
        OUTPUT_EXCEL,
        engine="openpyxl",
    ) as writer:

        summary_df.to_excel(
            writer,
            sheet_name="portfolio_summary",
            index=False,
        )

        counterparty_results.to_excel(
            writer,
            sheet_name="counterparty_results",
            index=False,
        )

        settings_df.to_excel(
            writer,
            sheet_name="model_settings",
            index=False,
        )

        model_notes.to_excel(
            writer,
            sheet_name="model_notes",
            index=False,
        )

        # Frozen fallback market data only.
        #
        # The as_of_date in this sheet refers to the fallback
        # snapshot and is not the requested market_data_date.
        market_data.to_excel(
            writer,
            sheet_name="fallback_market_data",
            index=False,
        )

        factor_proxies.to_excel(
            writer,
            sheet_name="factor_proxies",
            index=False,
        )

        rating_migration_matrix.to_excel(
            writer,
            sheet_name="rating_migration_matrix",
            index=False,
        )

        scenario_losses.to_excel(
            writer,
            sheet_name="scenario_losses",
            index=False,
        )

    format_excel_output()


# Remove temporary plot files when separate plot output is disabled.
def remove_plot_files():

    for plot_file in PLOT_FILES:

        if plot_file.exists():

            plot_file.unlink()


def main():

    OUTPUT_DIR.mkdir(
        exist_ok=True
    )

    # Run the complete model once.
    results = (
        run_credit_portfolio_model(
            FILE_PATH
        )
    )

    portfolio = results[
        "portfolio"
    ]

    settings = results[
        "settings"
    ]

    factor_proxies = results[
        "factor_proxies"
    ]

    rating_migration_matrix = (
        results[
            "rating_migration_matrix"
        ]
    )

    market_data = results[
        "market_data"
    ]

    portfolio_losses = results[
        "portfolio_losses"
    ]

    summary = results[
        "summary"
    ]

    # Read output choices from the central steering workbook.
    create_excel = get_output_flag(
        settings,
        "create_excel_output",
    )

    create_pdf = get_output_flag(
        settings,
        "create_pdf_output",
    )

    create_plot_files = (
        get_output_flag(
            settings,
            "create_plot_files",
        )
    )

    print(
        f"\nCreate Excel output: "
        f"{create_excel}"
    )

    print(
        f"Create PDF output: "
        f"{create_pdf}"
    )

    print(
        f"Create plot files: "
        f"{create_plot_files}"
    )

    # Plots are required internally for Excel and PDF output.
    plots_required = (
        create_excel
        or create_pdf
        or create_plot_files
    )

    if plots_required:

        create_plots(
            portfolio_losses,
            portfolio,
            summary,
        )

    # Detailed Excel output.
    if create_excel:

        create_excel_output(
            portfolio,
            settings,
            factor_proxies,
            rating_migration_matrix,
            market_data,
            portfolio_losses,
            summary,
        )

        print(
            f"\nExcel output created: "
            f"{OUTPUT_EXCEL}"
        )

    # Compact PDF report.
    if create_pdf:

        create_credit_portfolio_report(
            results=results,
            output_pdf=OUTPUT_PDF,
            loss_distribution_plot=(
                LOSS_DISTRIBUTION_PLOT
            ),
            tail_distribution_plot=(
                TAIL_DISTRIBUTION_PLOT
            ),
        )

        print(
            f"PDF output created: "
            f"{OUTPUT_PDF}"
        )

    # Keep PNG files only when explicitly requested.
    if create_plot_files:

        for plot_file in PLOT_FILES:

            print(
                f"Plot created: "
                f"{plot_file}"
            )

    else:

        remove_plot_files()


if __name__ == "__main__":

    main()